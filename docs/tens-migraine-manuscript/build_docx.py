# -*- coding: utf-8 -*-
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from content import BLOCKS

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")

def add_runs(paragraph, text, base_size=11):
    """Add runs to a paragraph, honoring **bold** markers."""
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            r.font.size = Pt(base_size)
        r = paragraph.add_run(m.group(1))
        r.bold = True
        r.font.size = Pt(base_size)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        r.font.size = Pt(base_size)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def build():
    doc = Document()

    # Base style
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    for section in doc.sections:
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

    for block in BLOCKS:
        kind = block[0]

        if kind == 'title':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(block[1])
            r.bold = True
            r.font.size = Pt(16)
            doc.add_paragraph()

        elif kind == 'h1':
            doc.add_heading(block[1], level=1)

        elif kind == 'h2':
            doc.add_heading(block[1], level=2)

        elif kind == 'h3':
            doc.add_heading(block[1], level=3)

        elif kind == 'p':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.15
            add_runs(p, block[1])

        elif kind == 'mono':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_after = Pt(8)
            for i, line in enumerate(block[1].split('\n')):
                r = p.add_run(line)
                r.font.name = 'Consolas'
                r.font.size = Pt(10)
                if i < len(block[1].split('\n')) - 1:
                    p.add_run().add_break()

        elif kind == 'bullet':
            for item in block[1]:
                p = doc.add_paragraph(style='List Bullet')
                add_runs(p, item)

        elif kind == 'numlist':
            for item in block[1]:
                p = doc.add_paragraph(style='List Number')
                add_runs(p, item)

        elif kind == 'table':
            caption, headers, rows = block[1], block[2], block[3]
            cap = doc.add_paragraph()
            cap.paragraph_format.space_before = Pt(10)
            cap.paragraph_format.space_after = Pt(4)
            r = cap.add_run(caption)
            r.bold = True
            r.font.size = Pt(9.5)

            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True

            hdr_cells = table.rows[0].cells
            for i, h in enumerate(headers):
                hdr_cells[i].text = ''
                p = hdr_cells[i].paragraphs[0]
                r = p.add_run(h)
                r.bold = True
                r.font.size = Pt(9)
                set_cell_shading(hdr_cells[i], 'D9E2F3')

            for row in rows:
                cells = table.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = ''
                    p = cells[i].paragraphs[0]
                    r = p.add_run(str(val))
                    r.font.size = Pt(9)

            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        elif kind == 'reflist':
            for i, ref in enumerate(block[1], start=1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                p.paragraph_format.space_after = Pt(4)
                r = p.add_run(f"{i}. ")
                r.font.size = Pt(10)
                r2 = p.add_run(ref)
                r2.font.size = Pt(10)

    out_path = '/workspace/paper_output/TENS_Migraine_Manuscript_Final.docx'
    doc.save(out_path)
    print('Saved:', out_path)

if __name__ == '__main__':
    build()
