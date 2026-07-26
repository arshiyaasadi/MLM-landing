from pathlib import Path
from html import escape
import re
import unicodedata

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


SOURCE = Path("REVISED_MIGRAINE_MANUSCRIPT.md")
OUTPUT = Path("REVISED_MIGRAINE_MANUSCRIPT_IPHONE.docx")
PDF_OUTPUT = Path("REVISED_MIGRAINE_MANUSCRIPT_IPHONE.pdf")


def clean_markdown(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("*", "").replace("`", "")
    return text.strip()


def add_text(paragraph, text: str) -> None:
    paragraph.add_run(clean_markdown(text))


document = Document()
section = document.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = document.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for style_name, size in (("Title", 17), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True

lines = SOURCE.read_text(encoding="utf-8").splitlines()
paragraph_buffer: list[str] = []


def flush_paragraph() -> None:
    if not paragraph_buffer:
        return
    paragraph = document.add_paragraph()
    add_text(paragraph, " ".join(paragraph_buffer))
    paragraph_buffer.clear()


for raw_line in lines:
    line = raw_line.strip()

    if not line:
        flush_paragraph()
        continue

    if line == "---":
        flush_paragraph()
        continue

    heading_match = re.match(r"^(#{1,3})\s+(.+)$", line)
    if heading_match:
        flush_paragraph()
        level = len(heading_match.group(1))
        text = clean_markdown(heading_match.group(2))
        if level == 1 and not document.paragraphs:
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(text)
        else:
            document.add_heading(text, level=level)
        continue

    if re.match(r"^[-*]\s+", line):
        flush_paragraph()
        paragraph = document.add_paragraph(style="List Bullet")
        add_text(paragraph, re.sub(r"^[-*]\s+", "", line))
        continue

    if re.match(r"^\d+\.\s+", line):
        flush_paragraph()
        paragraph = document.add_paragraph(style="List Number")
        add_text(paragraph, re.sub(r"^\d+\.\s+", "", line))
        continue

    paragraph_buffer.append(line)

flush_paragraph()

document.core_properties.title = (
    "Supraorbital versus infraorbital external trigeminal nerve stimulation "
    "for migraine prevention"
)
document.core_properties.subject = "Revised randomized trial manuscript"
document.core_properties.language = "en-US"
document.save(OUTPUT)


def pdf_safe(text: str) -> str:
    replacements = {
        "μ": "micro",
        "≥": ">=",
        "≤": "<=",
        "−": "-",
        "×": "x",
        "–": "-",
        "—": "-",
        "’": "'",
        "“": '"',
        "”": '"',
    }
    for original, replacement in replacements.items():
        text = text.replace(original, replacement)
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")


sample_styles = getSampleStyleSheet()
body_style = ParagraphStyle(
    "ManuscriptBody",
    parent=sample_styles["BodyText"],
    fontName="Helvetica",
    fontSize=9.5,
    leading=12.5,
    spaceAfter=6,
)
title_style = ParagraphStyle(
    "ManuscriptTitle",
    parent=sample_styles["Title"],
    fontName="Helvetica-Bold",
    fontSize=16,
    leading=20,
    alignment=TA_CENTER,
    spaceAfter=16,
)
heading_styles = {
    1: ParagraphStyle(
        "ManuscriptHeading1",
        parent=sample_styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=16,
        spaceBefore=10,
        spaceAfter=6,
    ),
    2: ParagraphStyle(
        "ManuscriptHeading2",
        parent=sample_styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11.5,
        leading=14,
        spaceBefore=8,
        spaceAfter=5,
    ),
    3: ParagraphStyle(
        "ManuscriptHeading3",
        parent=sample_styles["Heading3"],
        fontName="Helvetica-Bold",
        fontSize=10.5,
        leading=13,
        spaceBefore=6,
        spaceAfter=4,
    ),
}
bullet_style = ParagraphStyle(
    "ManuscriptBullet",
    parent=body_style,
    leftIndent=16,
    firstLineIndent=-8,
    bulletIndent=5,
)

story = []
pdf_paragraph_buffer: list[str] = []
is_first_heading = True


def flush_pdf_paragraph() -> None:
    if not pdf_paragraph_buffer:
        return
    text = pdf_safe(clean_markdown(" ".join(pdf_paragraph_buffer)))
    story.append(Paragraph(escape(text), body_style))
    pdf_paragraph_buffer.clear()


for raw_line in lines:
    line = raw_line.strip()
    if not line or line == "---":
        flush_pdf_paragraph()
        continue

    heading_match = re.match(r"^(#{1,3})\s+(.+)$", line)
    if heading_match:
        flush_pdf_paragraph()
        level = len(heading_match.group(1))
        text = escape(pdf_safe(clean_markdown(heading_match.group(2))))
        if is_first_heading:
            story.append(Paragraph(text, title_style))
            is_first_heading = False
        else:
            story.append(Paragraph(text, heading_styles[level]))
        continue

    if re.match(r"^[-*]\s+", line):
        flush_pdf_paragraph()
        text = re.sub(r"^[-*]\s+", "", line)
        text = escape(pdf_safe(clean_markdown(text)))
        story.append(Paragraph(text, bullet_style, bulletText="-"))
        continue

    if re.match(r"^\d+\.\s+", line):
        flush_pdf_paragraph()
        number, text = line.split(".", 1)
        text = escape(pdf_safe(clean_markdown(text)))
        story.append(Paragraph(text, bullet_style, bulletText=f"{number}."))
        continue

    pdf_paragraph_buffer.append(line)

flush_pdf_paragraph()


def add_page_number(canvas, pdf_document) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.drawCentredString(A4[0] / 2, 0.45 * inch, str(pdf_document.page))
    canvas.restoreState()


pdf_document = SimpleDocTemplate(
    str(PDF_OUTPUT),
    pagesize=A4,
    rightMargin=0.7 * inch,
    leftMargin=0.7 * inch,
    topMargin=0.7 * inch,
    bottomMargin=0.7 * inch,
    title=document.core_properties.title,
    subject=document.core_properties.subject,
)
pdf_document.build(story, onFirstPage=add_page_number, onLaterPages=add_page_number)

print(OUTPUT)
print(PDF_OUTPUT)
